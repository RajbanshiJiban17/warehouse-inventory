import io
from datetime import datetime
from decimal import Decimal
from typing import Any, List
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def export_as_csv(filename: str, headers: List[str], rows: List[List[Any]]) -> StreamingResponse:
    import csv
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([str(v) if v is not None else "" for v in row])
    output.seek(0)
    
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
    )


def export_as_excel(filename: str, title: str, headers: List[str], rows: List[List[Any]]) -> StreamingResponse:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]  # Excel worksheet title max 31 chars

    # Header style
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, row in enumerate(rows, start=2):
        formatted_row = []
        for val in row:
            if isinstance(val, Decimal):
                formatted_row.append(float(val))
            elif isinstance(val, datetime):
                formatted_row.append(val.strftime("%Y-%m-%d %H:%M"))
            else:
                formatted_row.append(val)
        ws.append(formatted_row)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=col_num)
            cell.font = data_font
            cell.border = thin_border

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    excel_io = io.BytesIO()
    wb.save(excel_io)
    excel_io.seek(0)

    return StreamingResponse(
        excel_io,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
    )


def export_as_pdf(filename: str, title: str, headers: List[str], rows: List[List[Any]]) -> StreamingResponse:
    buffer = io.BytesIO()
    # Landscape orientation for wide inventory tables
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=30,
        rightMargin=30,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        name="TitleStyle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A"),
        alignment=1,  # Centered
    )
    meta_style = ParagraphStyle(
        name="MetaStyle",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
        alignment=1,
    )

    elements = []
    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Warehouse Inventory System", meta_style))
    elements.append(Spacer(1, 15))

    # Convert row values to strings or short paragraphs
    cell_style = ParagraphStyle(name="Cell", parent=styles["Normal"], fontSize=8, leading=10)
    header_style = ParagraphStyle(name="HCell", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.white, fontName="Helvetica-Bold")

    table_data = [[Paragraph(str(h), header_style) for h in headers]]
    for row in rows:
        row_cells = []
        for val in row:
            if isinstance(val, Decimal):
                txt = f"{val:.2f}"
            elif isinstance(val, datetime):
                txt = val.strftime("%Y-%m-%d %H:%M")
            else:
                txt = str(val) if val is not None else ""
            row_cells.append(Paragraph(txt, cell_style))
        table_data.append(row_cells)

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))

    elements.append(t)
    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}.pdf"},
    )
